import unittest
import os
import time
import tempfile
from pathlib import Path

from src.db import init_db, upsert_job, get_unscored_jobs, update_job_score, get_jobs_for_application, get_summary_stats
from src.utils import normalize_job_url, load_settings, load_resume
from src.ranker import parse_json_from_response


class TestNaukriAgent(unittest.TestCase):
    def test_url_normalization(self):
        raw = "https://www.naukri.com/job-listings-python-dev-12345?src=search&sid=abc&xp=1"
        clean = normalize_job_url(raw)
        self.assertEqual(clean, "https://www.naukri.com/job-listings-python-dev-12345")

    def test_settings_and_resume_loading(self):
        settings = load_settings()
        self.assertIn("search", settings)
        self.assertIn("filters", settings)
        resume = load_resume()
        self.assertTrue(len(resume) > 50)

    def test_json_parsing(self):
        sample_markdown = "```json\n{\"score\": 88, \"justification\": \"Great match for Python\"}\n```"
        parsed = parse_json_from_response(sample_markdown)
        self.assertEqual(parsed["score"], 88)
        self.assertEqual(parsed["justification"], "Great match for Python")

    def test_db_lifecycle(self):
        init_db()
        test_url = f"https://www.naukri.com/job-listings-test-sample-job-{int(time.time()*1000)}"
        job_data = {
            "url": test_url,
            "title": "Senior Python Developer",
            "company": "Tech Corp",
            "location": "Bengaluru",
            "experience": "1-2 Yrs",
            "skills": "Python, FastAPI, Docker",
            "posted_date": "1 day ago",
            "description": "Looking for backend engineer with 1-2 years experience"
        }
        
        # Test insert
        inserted = upsert_job(job_data)
        # Deduplication test
        duplicate = upsert_job(job_data)
        self.assertFalse(duplicate)

        # Query unscored
        unscored = get_unscored_jobs(limit=500)
        matching = [j for j in unscored if j["url"] == job_data["url"]]
        self.assertTrue(len(matching) >= 1)
        job_id = matching[0]["id"]

        # Score job
        update_job_score(job_id, 85, "Matches Python experience well", "scored")
        
        # Verify eligible for apply
        eligible = get_jobs_for_application(min_score=75)
        self.assertTrue(any(j["id"] == job_id for j in eligible))

        # Check summary stats
        stats = get_summary_stats()
        self.assertIn("total", stats)

    def test_platform_registry_and_adapters(self):
        from src.platforms import get_platform, get_enabled_platforms, PLATFORM_REGISTRY
        from src.platforms.base import JobPlatform
        
        # Test registry contains naukri
        self.assertIn("naukri", PLATFORM_REGISTRY)
        naukri_adapter = get_platform("naukri")
        self.assertIsNotNone(naukri_adapter)
        self.assertIsInstance(naukri_adapter, JobPlatform)
        self.assertEqual(naukri_adapter.name, "naukri")

        # Test platform retrieval via settings
        settings = {"sites": ["naukri", "nonexistent"]}
        enabled = get_enabled_platforms(settings)
        self.assertEqual(len(enabled), 1)
        self.assertEqual(enabled[0].name, "naukri")

        # Test CLI override
        cli_enabled = get_enabled_platforms(settings, requested_sites=["naukri"])
        self.assertEqual(len(cli_enabled), 1)
        self.assertEqual(cli_enabled[0].name, "naukri")

    def test_platform_auth_and_sessions(self):
        from src.auth import (
            PLATFORM_AUTH_CONFIG, 
            get_platform_session_file, 
            get_platform_profile_dir,
            has_platform_session
        )
        
        expected_sites = ["naukri", "linkedin", "foundit", "indeed", "shine", "instahyre"]
        for site in expected_sites:
            self.assertIn(site, PLATFORM_AUTH_CONFIG)
            session_file = get_platform_session_file(site)
            self.assertTrue(str(session_file).endswith(f"{site}.json"))
            profile_dir = get_platform_profile_dir(site)
            self.assertTrue(profile_dir.exists())

    def test_phase3_excel_and_deduplication(self):
        import openpyxl
        from src.exporter import (
            EXCEL_HEADERS, 
            record_application_in_excel, 
            export_jobs_to_excel, 
            normalize_status,
            STATUS_APPLIED,
            STATUS_EXTERNAL_REDIRECT
        )
        from src.db import is_job_already_applied, upsert_job, update_job_apply_status

        # 1. Verify exact 8 required headers
        expected_headers = ["Date", "Site", "Company", "Job Title", "Location", "Job URL", "Status", "Reason"]
        self.assertEqual(EXCEL_HEADERS, expected_headers)

        # 2. Test status normalization
        self.assertEqual(normalize_status("applied"), STATUS_APPLIED)
        self.assertEqual(normalize_status("manual_needed"), STATUS_EXTERNAL_REDIRECT)
        self.assertEqual(normalize_status("external redirect"), STATUS_EXTERNAL_REDIRECT)

        # 3. Test atomic Excel writing and deduplication
        temp_excel = Path(tempfile.gettempdir()) / f"test_applied_{int(time.time()*1000)}.xlsx"
        sample_job = {
            "id": 9999,
            "platform": "linkedin",
            "title": "Lead AI Engineer",
            "company": "OpenAI",
            "location": "Remote",
            "url": "https://www.linkedin.com/jobs/view/999999"
        }

        # Write first time
        record_application_in_excel(sample_job, status="applied", reason="Passed all criteria", file_path=temp_excel)
        self.assertTrue(temp_excel.exists())

        wb = openpyxl.load_workbook(str(temp_excel))
        ws = wb.active
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(row=2, column=2).value, "Linkedin")
        self.assertEqual(ws.cell(row=2, column=7).value, "Applied")

        # Write duplicate update on same Site + URL -> Row count should stay 2, status updated
        record_application_in_excel(sample_job, status="already applied", reason="Updated status", file_path=temp_excel)
        wb2 = openpyxl.load_workbook(str(temp_excel))
        ws2 = wb2.active
        self.assertEqual(ws2.max_row, 2)
        wb.close()
        wb2.close()

        # Cleanup
        if temp_excel.exists():
            temp_excel.unlink()

    def test_linkedin_platform_integration(self):
        from src.platforms.linkedin import construct_linkedin_search_url, LinkedInPlatform
        from src.platforms import get_platform

        # 1. Test search URL construction with Easy Apply and filters
        url = construct_linkedin_search_url("GenAI Engineer", "Bengaluru", experience_years=3, job_age_days=30)
        self.assertIn("keywords=GenAI+Engineer", url)
        self.assertIn("location=Bengaluru", url)
        self.assertIn("f_AL=true", url)
        self.assertIn("f_E=", url)

        # 2. Test adapter instantiation
        adapter = get_platform("linkedin")
        self.assertIsNotNone(adapter)
        self.assertIsInstance(adapter, LinkedInPlatform)
        self.assertEqual(adapter.name, "linkedin")

    def test_experience_and_city_filters(self):
        from src.filters import (
            parse_experience_range,
            is_experience_matching,
            is_location_matching,
            filter_job_card,
            build_naukri_search_url,
            build_linkedin_search_url,
            build_indeed_search_url,
            build_foundit_search_url,
            build_shine_search_url,
            build_instahyre_search_url
        )

        # 1. Test experience string parsing
        self.assertEqual(parse_experience_range("3-5 Yrs"), (3.0, 5.0))
        self.assertEqual(parse_experience_range("0 - 2 years"), (0.0, 2.0))
        self.assertEqual(parse_experience_range("5+ Yrs"), (5.0, 99.0))
        self.assertEqual(parse_experience_range("3 Yrs"), (3.0, 3.0))
        self.assertEqual(parse_experience_range("Fresher"), (0.0, 1.0))

        # 2. Test experience matching logic (Target: 3 Yrs)
        self.assertTrue(is_experience_matching("2-5 Yrs", target_exp_years=3.0))
        self.assertTrue(is_experience_matching("3-6 Yrs", target_exp_years=3.0))
        self.assertTrue(is_experience_matching("0-2 Yrs", target_exp_years=3.0))
        # 8-15 Yrs should fail for a 3 Yrs candidate
        self.assertFalse(is_experience_matching("8-15 Yrs", target_exp_years=3.0, max_tolerance_years=6.0))
        self.assertFalse(is_experience_matching("10+ Yrs", target_exp_years=3.0, max_tolerance_years=6.0))

        # 3. Test city / location alias matching
        preferred = ["Bengaluru", "Remote", "Hyderabad"]
        self.assertTrue(is_location_matching("Bengaluru / Bangalore", preferred))
        self.assertTrue(is_location_matching("Whitefield, Bangalore", preferred))
        self.assertTrue(is_location_matching("Hyderabad / Secunderabad", preferred))
        self.assertTrue(is_location_matching("Work from home / Remote", preferred))
        self.assertTrue(is_location_matching("Virtual / Anywhere in India", preferred))
        
        # Mismatched location should fail
        self.assertFalse(is_location_matching("Noida, Uttar Pradesh", preferred, strict_match=True))
        self.assertFalse(is_location_matching("Pune, Maharashtra", preferred, strict_match=True))

        # 4. Test filter_job_card
        settings = {
            "search": {
                "experience_years": 3,
                "locations": ["Bengaluru", "Remote"]
            },
            "filters": {
                "experience_min": 1,
                "experience_max": 6,
                "strict_location_match": True
            }
        }

        # Valid card
        valid_card = {
            "title": "GenAI / LLM Engineer",
            "company": "AI Labs",
            "location": "Bengaluru",
            "experience": "2-5 Yrs"
        }
        passed, reason = filter_job_card(valid_card, settings)
        self.assertTrue(passed)

        # Invalid experience card
        high_exp_card = {
            "title": "Principal Architect",
            "company": "Big Corp",
            "location": "Bengaluru",
            "experience": "10-15 Yrs"
        }
        passed, reason = filter_job_card(high_exp_card, settings)
        self.assertFalse(passed)

        # Invalid location card
        wrong_loc_card = {
            "title": "AI Engineer",
            "company": "Startup",
            "location": "Kolkata, West Bengal",
            "experience": "3 Yrs"
        }
        passed, reason = filter_job_card(wrong_loc_card, settings)
        self.assertFalse(passed)

        # 5. Test search URL builders across platforms
        naukri_url = build_naukri_search_url("GenAI Engineer", "Bengaluru", experience_years=3, job_age_days=30)
        self.assertIn("experience=3", naukri_url)
        self.assertIn("jobAge=30", naukri_url)
        self.assertIn("bengaluru", naukri_url)

        indeed_url = build_indeed_search_url("GenAI Engineer", "Bengaluru", experience_years=3)
        self.assertIn("in.indeed.com", indeed_url)
        self.assertIn("explvl=MID_LEVEL", indeed_url)

        foundit_url = build_foundit_search_url("GenAI Engineer", "Bengaluru", experience_years=3)
        self.assertIn("foundit.in", foundit_url)
        self.assertIn("experience=3", foundit_url)

        shine_url = build_shine_search_url("GenAI Engineer", "Bengaluru", experience_years=3)
        self.assertIn("shine.com", shine_url)
        self.assertIn("exp=3", shine_url)

        instahyre_url = build_instahyre_search_url("GenAI Engineer", "Bengaluru", experience_years=3)
        self.assertIn("instahyre.com", instahyre_url)
        self.assertIn("years_of_experience=3", instahyre_url)

    def test_all_platform_adapters(self):
        from src.platforms import get_platform, PLATFORM_REGISTRY
        from src.platforms.naukri import NaukriPlatform
        from src.platforms.linkedin import LinkedInPlatform
        from src.platforms.foundit import FounditPlatform
        from src.platforms.indeed import IndeedPlatform
        from src.platforms.shine import ShinePlatform
        from src.platforms.instahyre import InstahyrePlatform

        expected_platforms = {
            "naukri": NaukriPlatform,
            "linkedin": LinkedInPlatform,
            "foundit": FounditPlatform,
            "indeed": IndeedPlatform,
            "shine": ShinePlatform,
            "instahyre": InstahyrePlatform
        }

        for site_name, site_class in expected_platforms.items():
            self.assertIn(site_name, PLATFORM_REGISTRY)
            adapter = get_platform(site_name)
            self.assertIsNotNone(adapter)
            self.assertIsInstance(adapter, site_class)
            self.assertEqual(adapter.name, site_name)
            self.assertTrue(len(adapter.display_name) > 0)


    def test_balanced_unscored_jobs_retrieval(self):
        # Insert test jobs for different platforms
        p1_job = {
            "url": f"https://www.linkedin.com/jobs/view/test-{int(time.time()*1000)}",
            "title": "LinkedIn AI Engineer",
            "company": "TestLinkedIn Corp",
            "location": "Bengaluru",
            "platform": "linkedin"
        }
        p2_job = {
            "url": f"https://www.shine.com/jobs/test-{int(time.time()*1000)}",
            "title": "Shine AI Engineer",
            "company": "TestShine Corp",
            "location": "Bengaluru",
            "platform": "shine"
        }
        upsert_job(p1_job)
        upsert_job(p2_job)

        # Test platform-specific retrieval
        linkedin_unscored = get_unscored_jobs(limit=10, platform="linkedin")
        self.assertTrue(all(j["platform"] == "linkedin" for j in linkedin_unscored))

        shine_unscored = get_unscored_jobs(limit=10, platform="shine")
        self.assertTrue(all(j["platform"] == "shine" for j in shine_unscored))

        # Test multi-platform balanced retrieval includes multiple platforms
        balanced = get_unscored_jobs(limit=15)
        platforms_found = {j["platform"] for j in balanced}
        self.assertTrue(len(platforms_found) >= 2)


if __name__ == "__main__":
    unittest.main()






